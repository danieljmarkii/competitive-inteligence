#!/usr/bin/env python3
"""Build the one-slide competitive landscape PPTX for a squad alignment deck.

Usage
  python3 tools/landscape_slide.py reports/2026-09-landscape-slide.json reports/2026-09-landscape-slide.pptx

Content JSON (spec: CI-Prompt.md -> Output 3). Every claim must trace to a
verified item in that month's report; no CONFIDENTIAL data.
  {
    "kicker":   "Competitive landscape · September 2026",
    "headline": "One sentence: the month's single most important shift",
    "stats":    [{"value": "31", "label": "verified updates"}, ...],      # 0-3 items
    "themes":   [{"title": "...", "evidence": "...", "so_what": "..."}],  # 3-4 items
    "watch_next": "...",
    "source":   "Source line: window, counts, report path",
    "notes":    "Speaker notes: talk track + source URLs"
  }

16:9 (13.333 x 7.5 in), Calibri throughout, white background. Requires python-pptx
(pip install python-pptx). Render to check fit before presenting.
"""

import json
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1B, 0x26, 0x44)   # dominant: headline, titles, numbers
INK = RGBColor(0x3A, 0x45, 0x5C)    # body text
TEAL = RGBColor(0x00, 0x7C, 0x7A)   # single accent: kicker, "so what" labels
MUTED = RGBColor(0x6B, 0x74, 0x86)  # captions, source line
CARD = RGBColor(0xF1, 0xF4, 0xF8)   # cool card tint
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

W, H, MARGIN = 13.333, 7.5, 0.6


def text_box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def add_run(paragraph, text, size, color, bold=False, italic=False, spacing=None):
    run = paragraph.add_run()
    run.text = text
    f = run.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = color
    if spacing:  # character spacing in 1/100 pt
        run._r.get_or_add_rPr().set("spc", str(spacing))
    return run


def build(content, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    stats = content.get("stats", [])[:3]
    step = 1.25  # stat column pitch; labels wrap to two lines inside it
    head_w = W - 2 * MARGIN - (step * len(stats) + 0.2 if stats else 0)

    # Kicker + headline
    p = text_box(slide, MARGIN, 0.45, head_w, 0.3).paragraphs[0]
    add_run(p, content["kicker"].upper(), 12, TEAL, bold=True, spacing=120)
    p = text_box(slide, MARGIN, 0.78, head_w, 1.2).paragraphs[0]
    p.line_spacing = 0.95
    add_run(p, content["headline"], 26, NAVY, bold=True)

    # Stat callouts, top right
    for i, s in enumerate(stats):
        x = W - MARGIN - (len(stats) - i) * step
        tf = text_box(slide, x, 0.55, step - 0.1, 0.6, MSO_ANCHOR.BOTTOM)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_run(p, s["value"], 30, NAVY, bold=True)
        tf = text_box(slide, x, 1.18, step - 0.1, 0.6)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_run(p, s["label"], 10.5, MUTED)

    # 2x2 theme cards
    themes = content["themes"][:4]
    card_w, card_h, gap = (W - 2 * MARGIN - 0.2) / 2, 1.95, 0.15
    for i, t in enumerate(themes):
        cx = MARGIN + (i % 2) * (card_w + 0.2)
        cy = 2.05 + (i // 2) * (card_h + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(cy),
                                      Inches(card_w), Inches(card_h))
        card.adjustments[0] = 0.07
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        card.shadow.inherit = False
        card.name = f"Theme card {i + 1}"

        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + 0.22), Inches(cy + 0.2),
                                     Inches(0.5), Inches(0.5))
        dot.fill.solid()
        dot.fill.fore_color.rgb = NAVY
        dot.line.fill.background()
        dot.shadow.inherit = False
        tf = dot.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_run(p, str(i + 1), 16, WHITE, bold=True)

        tx, tw = cx + 0.9, card_w - 1.15
        p = text_box(slide, tx, cy + 0.2, tw, 0.42).paragraphs[0]
        add_run(p, t["title"], 17, NAVY, bold=True)
        # Evidence and "so what" flow in one box so short evidence leaves no gap.
        tf = text_box(slide, tx, cy + 0.64, tw, card_h - 0.8)
        p = tf.paragraphs[0]
        p.line_spacing = 0.98
        add_run(p, t["evidence"], 13.5, INK)
        p = tf.add_paragraph()
        p.line_spacing = 0.98
        p.space_before = Pt(7)
        add_run(p, "So what for QW: ", 13.5, TEAL, bold=True)
        add_run(p, t["so_what"], 13.5, NAVY)

    # Watch next + source line
    if content.get("watch_next"):
        p = text_box(slide, MARGIN, 6.3, W - 2 * MARGIN, 0.4).paragraphs[0]
        add_run(p, "Watch next: ", 13, TEAL, bold=True)
        add_run(p, content["watch_next"], 13, INK)
    p = text_box(slide, MARGIN, 6.8, W - 2 * MARGIN, 0.3).paragraphs[0]
    add_run(p, content.get("source", ""), 10, MUTED)

    if content.get("notes"):
        slide.notes_slide.notes_text_frame.text = content["notes"]
    prs.core_properties.title = content["kicker"]
    prs.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as fh:
        build(json.load(fh), sys.argv[2])
    print(f"Wrote {sys.argv[2]}")
